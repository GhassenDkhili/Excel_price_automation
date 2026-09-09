import openpyxl as xl

wb = xl.load_workbook('transactions.xlsx')
sheet = wb['Feuil1']

for row in range(2, sheet.max_row + 1):
    cell = sheet.cell(row, column=3)

    price = float(cell.value.replace('$', ''))
    corrected_price = price * 0.9

    corrected_price_cell = sheet.cell(row, column=4)
    corrected_price_cell.value = corrected_price

wb.save('transactions_new.xlsx')